"""The branded PNM Queue B workbook.

This is the last mile: the point where a derived queue stops being a database object and
becomes a spreadsheet a person works through. Three things go wrong at the last mile, so
all three are controls here rather than conventions:

1. **Excel eats leading zeros.** Postcode ``05000`` becomes ``5000`` and state code ``00``
   becomes ``0`` unless the cell carries the Excel text format ``@``. That is the exact
   silent-relocation failure `grid.normalise.addresses.normalise_postcode` exists to
   prevent, reintroduced by the export. Every postcode, postcode-prefix and state column
   is written as text (`_TEXT_FORMAT`), and a test reads the file back to prove it.
2. **A stray workbook at the repo root fails the build.** Check 13 of
   ``scripts/check_context.py`` globs the filesystem for ``*.xlsx`` at the repo root, so
   being gitignored does not save you. `_assert_within_exports` refuses to write anywhere
   but a ``data/exports`` directory, and refuses the repo root explicitly.
3. **An export is the easiest place to leak personal data.** `QueueBRow` deliberately
   carries no telephone number, no practitioner name, no e-mail address and no tax
   identifier. `assert_business_only` holds that line from the other side: an explicit
   header allow-list plus a value scan for personal-data shapes, run *before* the file is
   saved so a leaky workbook never reaches disk.

Brand tokens come from ``docs/context/brand.md`` and are declared once at the top of this
module. Per that file this is **the** implementation of those tokens for workbooks — do
not re-declare fills, fonts or hexes at a call site.
"""

from __future__ import annotations

import datetime as dt
import re
import statistics
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, fields
from pathlib import Path
from types import MappingProxyType
from typing import Any, Final

import structlog
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from grid.panel.suppression import (
    QueueBPopulation,
    QueueBRow,
    Segment,
    SuppressionReport,
)
from grid.score.priority import Priority, PriorityBand, priority_key

log = structlog.get_logger(__name__)

# --------------------------------------------------------------------------------------
# Brand tokens — docs/context/brand.md and the `pmcare-brand` skill. Declared once.
# --------------------------------------------------------------------------------------

FONT_NAME: Final = "Arial"
"""Arial throughout, headings and body. No substitutes."""

NAVY: Final = "1F4E79"
"""Primary. Table header fills and the cover title block."""

TEAL: Final = "16A085"
"""Secondary. Informational notice bars."""

DARK_GRAY: Final = "2D3748"
"""Accent. Body emphasis, and notice text where white would not read."""

HEADER_BLUE: Final = "2C5282"
"""Header text on the cover's label column."""

STATUS_GREEN: Final = "10B981"
STATUS_AMBER: Final = "F59E0B"
STATUS_RED: Final = "EF4444"
STATUS_BLUE: Final = "3B82F6"
WHITE: Final = "FFFFFF"

BAND_FILLS: Final[Mapping[PriorityBand, str]] = MappingProxyType(
    {
        PriorityBand.A_CHAIN_FOOTHOLD: STATUS_GREEN,
        PriorityBand.B_RECENT_RECORD: STATUS_AMBER,
        PriorityBand.C_CONTACTABLE: STATUS_BLUE,
        PriorityBand.D_NEEDS_CONTACT: STATUS_RED,
    }
)
"""Band A/B/C/D to the four status colours: green, amber, blue, red."""

_LIGHT_TEXT_ON: Final[frozenset[str]] = frozenset({STATUS_AMBER, STATUS_GREEN})
"""Fills too light for white text. Dark Gray is used on these instead."""

# --------------------------------------------------------------------------------------
# Structure
# --------------------------------------------------------------------------------------

SHEET_COVER: Final = "Cover"
SHEET_QUEUE_B: Final = "Queue B"
SHEET_WIN_BACK: Final = "Win-back"
SHEET_REVIEW: Final = "Review"
SHEET_FUNNEL: Final = "Suppression funnel"
SHEET_NOTES: Final = "Notes & definitions"

SHEET_ORDER: Final[tuple[str, ...]] = (
    SHEET_COVER,
    SHEET_QUEUE_B,
    SHEET_WIN_BACK,
    SHEET_REVIEW,
    SHEET_FUNNEL,
    SHEET_NOTES,
)

NOTICE_ROW: Final = 1
"""Row 1 of every grid sheet carries a one-line notice, not data."""

HEADER_ROW: Final = 2
FIRST_DATA_ROW: Final = 3

_TEXT_FORMAT: Final = "@"
"""Excel's text format. The whole reason ``05000`` survives the round trip."""

_DATE_FORMAT: Final = "dd mmm yyyy"

_MONTHS: Final[tuple[str, ...]] = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)
"""British English month names, spelled out rather than taken from ``strftime('%B')``.
``%B`` follows the process locale, so on a machine with a non-English locale the brand's
``DD Month YYYY`` would silently render in another language."""

BAND_HEADER: Final = "Priority band"
"""The one column that carries a status fill, so it is named rather than positional."""

# --------------------------------------------------------------------------------------
# Standing notices
# --------------------------------------------------------------------------------------

DRAFT_NOTICE: Final = "DRAFT — requires human review before use or circulation"

NEUTRAL_TPA_NOTICE: Final = (
    "PMCare is a neutral third-party administrator (TPA/MCO). It is not the insurer and "
    "not a treating clinician. Nothing in this workbook is a clinical or underwriting "
    "decision."
)

BUSINESS_ONLY_NOTICE: Final = (
    "Business data only. No practitioner name, telephone number, e-mail address or tax "
    "identifier appears in this workbook (PDPA 2010 as amended 2024)."
)

PERSONAL_DATA_NOTICE: Final = (
    "This workbook contains personal data. Handle under the PMCare retention and access "
    "rules, do not forward it outside the authorised recipients, and delete your copy "
    "when the engagement it supports is closed."
)

QUEUE_B_NOTICE: Final = (
    "Queue B — registered GP clinics not on the PMCare panel, ranked. "
    "Every row already holds a PMCare provider record: this is re-engagement, not a "
    "first approach. " + DRAFT_NOTICE
)

WIN_BACK_NOTICE: Final = (
    "Win-back — a different conversation. These clinics were on the panel and were "
    "terminated; they are not never-engaged prospects. Ask why they left before pitching."
)

REVIEW_NOTICE: Final = (
    "Do not call from this sheet. Each row is either a probable duplicate of a clinic "
    "already on panel, or carries a PNM REMARKS flag. Resolve the reason first."
)

FUNNEL_NOTICE: Final = (
    "How the headline number was reached. Every step is counted so the figure you have "
    "been handed can be reconciled rather than trusted."
)

NOTES_NOTICE: Final = (
    "What each column means, and which values are known versus inferred. Read this before "
    "using chain grouping as an argument in a conversation."
)

# --------------------------------------------------------------------------------------
# The export's own PDPA control
# --------------------------------------------------------------------------------------

PERSONAL_DATA_PATTERNS: Final[Mapping[str, re.Pattern[str]]] = MappingProxyType(
    {
        "Malaysian NRIC": re.compile(r"\b\d{6}-\d{2}-\d{4}\b"),
        "Malaysian mobile number": re.compile(r"\+?60\s?1\d[\s-]?\d{3,4}[\s-]?\d{3,4}"),
        "Malaysian mobile number, local form": re.compile(r"\b01\d[\s-]?\d{3,4}[\s-]?\d{4}\b"),
        "e-mail address": re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"),
    }
)
"""Shapes that mean personal data has reached a business-only workbook.

Modelled on ``DISCLOSURE_PATTERNS`` in ``tests/unit/test_pr001_pdpa.py``. The scan is a
canary, not a classifier: it catches the shapes a leak actually takes, and it is deliberately
run over rendered cell values rather than over the source objects, because the failure mode
is somebody widening a column, not somebody widening a dataclass."""


class ExportPathError(RuntimeError):
    """A workbook was about to be written outside ``data/exports``.

    Deliberately fatal rather than a fallback-to-a-safe-path, because the two paths a
    caller most plausibly reaches for — the repo root and the current working directory —
    are the two that break the build (check 13) or commit real clinic data (guardrail 5).
    """


class PersonalDataInExportError(RuntimeError):
    """A business-only workbook was about to carry personal data.

    Raised before the file is saved, so the leaky workbook never exists on disk.
    """


# --------------------------------------------------------------------------------------
# Paths
# --------------------------------------------------------------------------------------

_REPO_ROOT: Final = Path(__file__).resolve().parents[3]
"""``src/grid/export/xlsx.py`` → repo root. Correct for this src-layout editable install;
callers who relocate the package should pass ``root`` explicitly."""

_DATA_DIR_NAME: Final = "data"
_EXPORTS_DIR_NAME: Final = "exports"


def exports_dir(root: Path | None = None) -> Path:
    """The directory every workbook is written into.

    Args:
        root: Directory to treat as the repository root. Injectable so tests can point the
            whole export path at ``tmp_path`` instead of the working tree.

    Returns:
        ``<root or repo>/data/exports``, resolved. Not created — `write_queue_b_workbook`
        creates it once the path guard has passed.
    """
    base = _REPO_ROOT if root is None else Path(root)
    return (base / _DATA_DIR_NAME / _EXPORTS_DIR_NAME).resolve()


def _assert_within_exports(path: Path, root: Path | None = None) -> Path:
    """Refuse to write a workbook anywhere but a ``data/exports`` directory.

    ``data/`` is gitignored in full (guardrail 5) and check 13 of ``check_context.py``
    fails the build on a stray ``*.xlsx`` at the repo root. Check 13 globs the filesystem
    rather than the git index, so a gitignored stray still fails it — which is why this is
    a hard error and not a warning.

    Args:
        path: The intended workbook path.
        root: When given, the only accepted parent is ``exports_dir(root)``. When omitted,
            any ``.../data/exports/`` directory is accepted, so that a caller writing under
            a temporary or relocated tree is not forced to thread a root through.

    Returns:
        The resolved path.

    Raises:
        ExportPathError: When the parent is the repo root, is not an exports directory, or
            when the file is not an ``.xlsx``.
    """
    resolved = Path(path).expanduser().resolve()
    parent = resolved.parent

    if resolved.suffix.lower() != ".xlsx":
        suffix = resolved.suffix or "no suffix"
        raise ExportPathError(
            f"{resolved.name}: a workbook must be written as .xlsx, not {suffix}."
        )

    # Checked before anything else and regardless of `root`: this is the specific mistake
    # that fails the build for everybody, not just the person who made it.
    if parent == _REPO_ROOT:
        raise ExportPathError(
            f"{resolved.name}: refusing to write a workbook to the repository root. "
            "Check 13 of scripts/check_context.py fails the build on any *.xlsx there, "
            f"and it globs the filesystem, so .gitignore does not help. Use "
            f"{exports_dir()} instead."
        )

    if root is not None:
        allowed = exports_dir(root)
        if parent != allowed:
            raise ExportPathError(
                f"{resolved}: exports must be written to {allowed}, not {parent}."
            )
        return resolved

    if parent == exports_dir():
        return resolved
    if parent.name == _EXPORTS_DIR_NAME and parent.parent.name == _DATA_DIR_NAME:
        return resolved

    raise ExportPathError(
        f"{resolved}: exports must be written to a data/exports directory "
        f"(for this working tree, {exports_dir()}). Real clinic data must never land "
        "outside the gitignored data/ tree (guardrail 5)."
    )


# --------------------------------------------------------------------------------------
# Workbook metadata
# --------------------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class WorkbookMeta:
    """Everything the cover sheet states about the workbook's provenance.

    Attributes:
        title: The workbook title, shown in the navy title block.
        as_of: The date the queue was evaluated on. A queue is always "as at" a date;
            leaving it implicit is how a stale list gets worked as a fresh one.
        generated_at: When the file was produced.
        source_note: The extract vintage the figures were derived from.
        contains_personal_data: False for the default business-only workbook. When False,
            `assert_business_only` is run over the whole workbook before it is saved.
        requested_by: Who asked for the extract, for the audit trail.
        purpose: What it is to be used for, per the PDPA purpose-limitation principle.
    """

    title: str
    as_of: dt.date
    generated_at: dt.datetime
    source_note: str
    contains_personal_data: bool
    requested_by: str | None = None
    purpose: str | None = None


# --------------------------------------------------------------------------------------
# Formatting helpers
# --------------------------------------------------------------------------------------


def _fill(colour: str) -> Any:
    """A solid fill in a brand colour."""
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


def _font(
    *, bold: bool = False, colour: str = DARK_GRAY, size: int = 10, italic: bool = False
) -> Any:
    """An Arial font in a brand colour."""
    return Font(name=FONT_NAME, bold=bold, color=colour, size=size, italic=italic)


def _text_colour_on(fill_colour: str) -> str:
    """Readable text colour for a given fill."""
    return DARK_GRAY if fill_colour in _LIGHT_TEXT_ON else WHITE


def format_date(value: dt.date) -> str:
    """Render a date as the brand's ``DD Month YYYY``.

    Args:
        value: The date to render.

    Returns:
        e.g. ``05 August 2026``. Zero-padded, British English month name.
    """
    return f"{value.day:02d} {_MONTHS[value.month - 1]} {value.year}"


def _format_datetime(value: dt.datetime) -> str:
    """Render a timestamp as ``DD Month YYYY at HH:MM``."""
    return f"{format_date(value.date())} at {value.hour:02d}:{value.minute:02d}"


def _yes_no(value: bool | None) -> str:
    """Render a nullable boolean without collapsing "unknown" into "no"."""
    if value is None:
        return "Unknown"
    return "Yes" if value else "No"


def _as_date(value: dt.datetime | None) -> dt.date | None:
    """Drop the time component, which PR001 never populates meaningfully."""
    return None if value is None else value.date()


# --------------------------------------------------------------------------------------
# Column specifications
# --------------------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _RowView:
    """One outlet as the workbook sees it: the row, its rank and its priority."""

    rank: int
    row: QueueBRow
    priority: Priority | None


@dataclass(frozen=True, slots=True)
class _Column:
    """One workbook column: its header, its width and how to get its value."""

    header: str
    width: int
    value: Callable[[_RowView], object]
    text_format: bool = False
    """True for columns Excel would otherwise read as numbers and strip zeros from."""


_SEGMENT_LABELS: Final[Mapping[Segment, str]] = MappingProxyType(
    {
        Segment.NOT_ON_PANEL: "Not on panel — call list",
        Segment.WIN_BACK: "Win-back — terminated",
        Segment.REVIEW_LIKELY_ON_PANEL: "Review — probably already on panel",
        Segment.REVIEW_FLAGGED: "Review — REMARKS flag",
    }
)


def _band_cell(view: _RowView) -> str:
    """The band column's value, including the unranked case.

    The band's own value is the label — `PriorityBand` members read as
    ``A — chain already on panel`` — so nothing is re-worded here. A second wording of the
    same idea is a second thing to keep in step, and it would drift.
    """
    if view.priority is None:
        return "Unranked"
    return view.priority.band.value


def _reasons_cell(view: _RowView) -> str:
    """`Priority.reasons`, rendered for a human reading one row at a time."""
    if view.priority is None or not view.priority.reasons:
        return "No ranking signal recorded."
    return "; ".join(view.priority.reasons)


def _review_reason(view: _RowView) -> str:
    """Why this row was routed to review rather than onto the call list."""
    row = view.row
    if row.segment is Segment.REVIEW_LIKELY_ON_PANEL:
        collides = row.collides_with_provider_code or "an on-panel outlet"
        return (
            f"Name and postcode collide with {collides}, which is already on panel. "
            "Probably the same clinic keyed twice. Confirm before contacting."
        )
    if row.segment is Segment.REVIEW_FLAGGED:
        flags = ", ".join(row.review_flags) or "an unspecified concern"
        return (
            f"PMCare's own REMARKS flag this provider: {flags}. Resolve the flag before contacting."
        )
    return "Routed for review. Confirm before contacting."


_IDENTITY_COLUMNS: Final[tuple[_Column, ...]] = (
    _Column("Provider code", 16, lambda v: v.row.provider_code),
    _Column("Clinic name", 42, lambda v: v.row.name_raw),
    _Column("Normalised name", 42, lambda v: v.row.name_normalised),
    _Column("Chain (inferred)", 26, lambda v: v.row.chain_base_name),
    _Column("Branch", 20, lambda v: v.row.branch_qualifier),
    _Column("Chain outlets on panel (inferred)", 18, lambda v: v.row.chain_on_panel_outlets),
    _Column("Unit", 16, lambda v: v.row.address_unit),
    _Column("Street", 36, lambda v: v.row.address_street),
    _Column("Locality", 28, lambda v: v.row.address_locality),
    _Column("Town", 22, lambda v: v.row.city_raw),
    _Column("Postcode", 12, lambda v: v.row.postcode_clean, text_format=True),
    _Column("Postcode valid", 14, lambda v: _yes_no(v.row.postcode_valid)),
    _Column("Postcode prefix", 14, lambda v: v.row.postcode_prefix, text_format=True),
    _Column("Postcode vs state mismatch", 22, lambda v: _yes_no(v.row.postcode_state_mismatch)),
    _Column("State code", 12, lambda v: v.row.state_code, text_format=True),
    _Column("Appointment date", 16, lambda v: _as_date(v.row.appointment_date)),
)
"""Business columns shared by all three list sheets. No personal-data column belongs here,
and `assert_business_only` is what stops one being added quietly."""

QUEUE_B_COLUMNS: Final[tuple[_Column, ...]] = (
    _Column("Rank", 7, lambda v: v.rank),
    _Column(BAND_HEADER, 24, _band_cell),
    _Column("Why this rank", 62, _reasons_cell),
    *_IDENTITY_COLUMNS,
)

WIN_BACK_COLUMNS: Final[tuple[_Column, ...]] = (
    *_IDENTITY_COLUMNS,
    _Column("Termination date", 16, lambda v: _as_date(v.row.termination_date)),
    _Column("Panel flag conflict", 18, lambda v: _yes_no(v.row.panel_flag_conflict)),
)

REVIEW_COLUMNS: Final[tuple[_Column, ...]] = (
    *_IDENTITY_COLUMNS,
    _Column("Review reason", 62, _review_reason),
    _Column("Collides with provider code", 24, lambda v: v.row.collides_with_provider_code),
    _Column("REMARKS flags", 28, lambda v: ", ".join(v.row.review_flags) or None),
    _Column("Segment", 32, lambda v: _SEGMENT_LABELS[v.row.segment]),
)

FUNNEL_HEADERS: Final[tuple[str, ...]] = ("Step", "Value", "What it means")
NOTES_HEADERS: Final[tuple[str, ...]] = ("Column or note", "What it means", "Known or inferred")

ALLOWED_HEADERS: Final[frozenset[str]] = frozenset(
    {
        # Queue B
        "Rank",
        BAND_HEADER,
        "Why this rank",
        # Shared identity and address block
        "Provider code",
        "Clinic name",
        "Normalised name",
        "Chain (inferred)",
        "Branch",
        "Chain outlets on panel (inferred)",
        "Unit",
        "Street",
        "Locality",
        "Town",
        "Postcode",
        "Postcode valid",
        "Postcode prefix",
        "Postcode vs state mismatch",
        "State code",
        "Appointment date",
        # Win-back
        "Termination date",
        "Panel flag conflict",
        # Review
        "Review reason",
        "Collides with provider code",
        "REMARKS flags",
        "Segment",
        # Funnel and notes
        "Step",
        "Value",
        "What it means",
        "Column or note",
        "Known or inferred",
    }
)
"""Every header a business-only workbook may carry, written out by hand.

Hand-written on purpose. Derived from the column specs it would auto-approve whatever was
added to them, which is precisely the change it exists to catch: a header like "Contact
number" has to be added *here* as well, and that is a line a reviewer will see."""


# --------------------------------------------------------------------------------------
# Documentation content
# --------------------------------------------------------------------------------------

_COLUMN_NOTES: Final[Mapping[str, tuple[str, str]]] = MappingProxyType(
    {
        "Rank": ("Position in the call order, best first. Recomputed each run.", "Derived"),
        BAND_HEADER: (
            "A to D. A is the strongest reason to call first; D means contact details must "
            "be found before the clinic can be approached at all.",
            "Derived",
        ),
        "Why this rank": (
            "The signals behind the band, in the order they were applied. If the reason "
            "does not survive contact with the clinic, the ranking was wrong, not the clinic.",
            "Derived",
        ),
        "Provider code": (
            "PMCare's own provider code from the incumbent master. The join key back to "
            "any other PMCare system.",
            "Known",
        ),
        "Clinic name": ("The clinic name exactly as it is held in the source.", "Known"),
        "Normalised name": (
            "The name after case, punctuation and abbreviation normalisation. Used for "
            "matching, not for addressing the clinic.",
            "Derived",
        ),
        "Chain (inferred)": (
            "The base name shared with other outlets. INFERRED: a shared base name can mean "
            "a real brand, or the same premises keyed twice under two codes.",
            "Inferred",
        ),
        "Branch": ("The branch qualifier separated off the name, where one was found.", "Derived"),
        "Chain outlets on panel (inferred)": (
            "How many outlets sharing this base name are already on the panel. A real "
            "foothold is a strong opener; an inferred one is an embarrassment. Verify first.",
            "Inferred",
        ),
        "Unit": ("Lot, unit or shoplot number parsed off the address.", "Derived"),
        "Street": ("Street line parsed off the address (Jalan, Lorong, Persiaran).", "Derived"),
        "Locality": ("Taman, Bandar or Kampung line parsed off the address.", "Derived"),
        "Town": (
            "Free-text town from staging. Not canonicalised — 1,019 distinct spellings are "
            "not a vocabulary. Useful for planning a route, not for grouping.",
            "Known",
        ),
        "Postcode": (
            "Five-digit Malaysian postcode, written as Excel text so the leading zero "
            "survives. Reformatting this column as a number turns 05000 into 5000 and moves "
            "the clinic to another state.",
            "Known",
        ),
        "Postcode valid": ("Whether the postcode is five digits after cleaning.", "Derived"),
        "Postcode prefix": (
            "First two digits, the state-level block. Also text, for the same reason.",
            "Derived",
        ),
        "Postcode vs state mismatch": (
            "Yes where the postcode's modal state disagrees with the recorded state code. "
            "Unknown where it could not be evaluated. Neither field is assumed correct.",
            "Derived",
        ),
        "State code": (
            "PMCare's two-character state code, written as text. Code 00 is a real value in "
            "the extract with no known meaning, and it is a zero if this column is numeric.",
            "Known",
        ),
        "Appointment date": (
            "When this provider was first appointed. Present on every row here, which is why "
            "this is a re-engagement list rather than a prospect list.",
            "Known",
        ),
        "Termination date": (
            "When the panel relationship ended. On the Win-back sheet this is the whole "
            "point of the row.",
            "Known",
        ),
        "Panel flag conflict": (
            "Yes where a row is terminated and still flagged as on-panel. The two fields "
            "contradict each other; a human has to decide which is true.",
            "Derived",
        ),
        "Review reason": (
            "Why the row was held back from the call list: a colliding provider code, or a "
            "PNM REMARKS flag.",
            "Derived",
        ),
        "Collides with provider code": (
            "The on-panel provider code this row shares a normalised name and postcode with.",
            "Derived",
        ),
        "REMARKS flags": (
            "Signals mined from PMCare's own free-text REMARKS: closed, known duplicate, "
            "duplicate flag, code supersession.",
            "Derived",
        ),
        "Segment": ("Which list the row was routed to and why.", "Derived"),
    }
)
"""Plain-English meaning and provenance for every data column in the workbook."""

_FUNNEL_NOTES: Final[Mapping[str, str]] = MappingProxyType(
    {
        "as_of": (
            "The date activity was evaluated on. A queue is always 'as at' a date; a run "
            "with a different date is a different queue."
        ),
        "total_outlets": "Every outlet row in the incumbent master, all provider types.",
        "gp_rows": (
            "GP and primary care only. Dental, specialist, physiotherapy, laboratory, "
            "hospital and aesthetics-only providers are out of scope for GRID."
        ),
        "gp_active_as_of": (
            "Active by the point-in-time rule: appointed, not yet terminated, not "
            "suspended, as at the date above."
        ),
        "gp_active_status_code_a": (
            "Active by status_code = 'A' alone. Shown so the two definitions can be "
            "compared rather than assumed equal."
        ),
        "status_code_disagreements": (
            "Rows where the two definitions above disagree. Future-dated terminations mean "
            "this number grows over time, which is why status_code is not used."
        ),
        "on_panel": (
            "Active GP outlets already on the PMCare panel. Suppressed from the call list "
            "by a filter on the authoritative panel flag, not by a fuzzy match."
        ),
        "not_on_panel": (
            "Active GP outlets not on the panel. The population the three list sheets are "
            "drawn from."
        ),
        "excluded_superseded": (
            "Provider codes PMCare's own notes say were replaced by another code. Excluded "
            "rather than routed, because the replacement code is the row to work."
        ),
        "routed_to_review_flagged": (
            "Held back to the Review sheet: REMARKS say closed, duplicate or superseded."
        ),
        "routed_to_review_likely_on_panel": (
            "Held back to the Review sheet: normalised name and postcode collide with an "
            "outlet already on the panel."
        ),
        "win_back": (
            "Terminated since the win-back cut-off. Counted separately and not part of the "
            "reconciliation below, because a terminated outlet is not in 'not on panel'."
        ),
        "queue_b_final": "The call list on the Queue B sheet. This is the headline number.",
    }
)


# --------------------------------------------------------------------------------------
# The business-only control
# --------------------------------------------------------------------------------------


def assert_business_only(headers: Sequence[str], cells: Iterable[object]) -> None:
    """Assert a workbook carries business data only.

    Two checks, because a leak arrives by two routes. A new column is the deliberate route:
    somebody widens the export to "just add the phone number", so every header is checked
    against `ALLOWED_HEADERS`. A leaked value is the accidental route: a telephone number
    typed into a name field upstream, or a free-text column carrying more than it should,
    so every rendered value is scanned for personal-data shapes.

    The message names the pattern and the count but never quotes the value, so a failure
    does not itself become a second disclosure.

    Args:
        headers: Every column header in the workbook.
        cells: Every value written to the workbook, in any order.

    Raises:
        PersonalDataInExportError: When a header is outside the allow-list, or a value
            matches a telephone, NRIC or e-mail shape.
    """
    findings: list[str] = []

    unexpected = sorted({header for header in headers if header not in ALLOWED_HEADERS})
    if unexpected:
        findings.append(
            f"  columns outside the business-only allow-list: {unexpected}. "
            "Add the column to ALLOWED_HEADERS only after confirming it is business data."
        )

    counts: dict[str, int] = {}
    for value in cells:
        if value is None:
            continue
        rendered = str(value)
        for label, pattern in PERSONAL_DATA_PATTERNS.items():
            if pattern.search(rendered):
                counts[label] = counts.get(label, 0) + 1
    findings.extend(f"  {label}: {count} cell(s)" for label, count in sorted(counts.items()))

    if findings:
        raise PersonalDataInExportError(
            "\n".join(
                [
                    "PDPA — this workbook was declared business-data-only but carries "
                    "personal data:",
                    *findings,
                    "QueueBRow deliberately holds no telephone number, practitioner name, "
                    "e-mail address or tax identifier. Fix the source of the value rather "
                    "than removing this check; a workbook that can contain personal data "
                    "must say so on its cover and name its retention expectation.",
                ]
            )
        )


# --------------------------------------------------------------------------------------
# Sheet writing
# --------------------------------------------------------------------------------------


def _write_notice(ws: Any, text: str, *, colour: str, columns: int) -> None:
    """Write the merged notice bar that opens every grid sheet."""
    ws.merge_cells(
        start_row=NOTICE_ROW, start_column=1, end_row=NOTICE_ROW, end_column=max(columns, 2)
    )
    cell = ws.cell(row=NOTICE_ROW, column=1, value=text)
    cell.fill = _fill(colour)
    cell.font = _font(bold=True, colour=_text_colour_on(colour))
    cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=False)
    ws.row_dimensions[NOTICE_ROW].height = 22


def _write_grid(
    ws: Any,
    *,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    widths: Sequence[int],
    notice: str,
    notice_colour: str,
    text_columns: frozenset[int] = frozenset(),
) -> list[object]:
    """Write a notice bar, a navy header row and the data beneath it.

    Args:
        ws: The worksheet.
        headers: Column headers, written to `HEADER_ROW`.
        rows: Data rows, written from `FIRST_DATA_ROW`.
        widths: One column width per header.
        notice: The one-line notice for row 1.
        notice_colour: Brand colour for the notice bar.
        text_columns: Zero-based indices to write in Excel text format, so that leading
            zeros survive. Postcode and state columns, always.

    Returns:
        Every value written, for `assert_business_only` to scan.
    """
    _write_notice(ws, notice, colour=notice_colour, columns=len(headers))

    header_fill = _fill(NAVY)
    header_font = _font(bold=True, colour=WHITE)
    body_font = _font()
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.column_dimensions[str(get_column_letter(index))].width = widths[index - 1]
    ws.row_dimensions[HEADER_ROW].height = 30

    written: list[object] = list(headers)
    for offset, row in enumerate(rows):
        excel_row = FIRST_DATA_ROW + offset
        for index, value in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=index, value=value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index - 1 in text_columns:
                # The whole reason 05000 does not become 5000 on the way out.
                cell.number_format = _TEXT_FORMAT
            elif isinstance(value, dt.date) and not isinstance(value, dt.datetime):
                cell.number_format = _DATE_FORMAT
            written.append(value)

    last_column = str(get_column_letter(max(len(headers), 1)))
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{HEADER_ROW + len(rows)}"
    return written


def _write_list_sheet(
    ws: Any,
    columns: Sequence[_Column],
    views: Sequence[_RowView],
    *,
    notice: str,
    notice_colour: str,
) -> tuple[list[str], list[object]]:
    """Write one of the three list sheets from a column specification.

    Returns:
        The headers and every value written.
    """
    headers = [column.header for column in columns]
    rows = [[column.value(view) for column in columns] for view in views]
    written = _write_grid(
        ws,
        headers=headers,
        rows=rows,
        widths=[column.width for column in columns],
        notice=notice,
        notice_colour=notice_colour,
        text_columns=frozenset(index for index, column in enumerate(columns) if column.text_format),
    )
    return headers, written


def _colour_band_column(ws: Any, columns: Sequence[_Column], views: Sequence[_RowView]) -> None:
    """Fill the band cell green, amber, blue or red for A, B, C and D."""
    try:
        position = next(
            index for index, column in enumerate(columns, start=1) if column.header == BAND_HEADER
        )
    except StopIteration:  # pragma: no cover - the band column is part of the spec
        return

    for offset, view in enumerate(views):
        if view.priority is None:
            continue
        colour = BAND_FILLS.get(view.priority.band)
        if colour is None:
            continue
        cell = ws.cell(row=FIRST_DATA_ROW + offset, column=position)
        cell.fill = _fill(colour)
        cell.font = _font(bold=True, colour=_text_colour_on(colour))


def _write_cover(ws: Any, meta: WorkbookMeta, report: SuppressionReport) -> list[object]:
    """Write the branded cover sheet.

    Carries the title block, the provenance a reader needs to trust the numbers, the draft
    notice and the neutral-TPA line. Both notices are mandatory on every outward artefact.

    Returns:
        Every value written, for `assert_business_only` to scan.
    """
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 96
    written: list[object] = []

    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=2)
    title = ws.cell(row=1, column=1, value=meta.title)
    title.fill = _fill(NAVY)
    title.font = _font(bold=True, colour=WHITE, size=18)
    title.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.cell(row=2, column=1).fill = _fill(NAVY)
    ws.cell(row=1, column=2).fill = _fill(NAVY)
    ws.cell(row=2, column=2).fill = _fill(NAVY)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    written.append(meta.title)

    subtitle = ws.cell(
        row=3,
        column=1,
        value="Project GRID · Provider Network Management · PMCare Sdn Bhd",
    )
    subtitle.font = _font(bold=True, colour=TEAL, size=11)
    written.append(subtitle.value)

    personal_data = (
        "Yes — handle under the PMCare retention and access rules"
        if meta.contains_personal_data
        else "No — business data only"
    )
    facts: tuple[tuple[str, str], ...] = (
        ("As at", format_date(meta.as_of)),
        ("Generated", _format_datetime(meta.generated_at)),
        ("Source", meta.source_note),
        ("Requested by", meta.requested_by or "Not recorded"),
        (
            "Purpose",
            meta.purpose or "Provider network administration and panel engagement. Not marketing.",
        ),
        ("Contains personal data", personal_data),
        (
            "Call list",
            f"{report.queue_b_final:,} clinics, from {report.not_on_panel:,} active GP "
            f"outlets not on the panel. See the Suppression funnel sheet.",
        ),
    )

    row = 5
    for label, value in facts:
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = _font(bold=True, colour=HEADER_BLUE)
        label_cell.alignment = Alignment(vertical="top")
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = _font()
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
        written.extend((label, value))
        row += 1

    row += 1
    notices: tuple[tuple[str, str], ...] = (
        (DRAFT_NOTICE, STATUS_AMBER),
        (NEUTRAL_TPA_NOTICE, DARK_GRAY),
        (
            PERSONAL_DATA_NOTICE if meta.contains_personal_data else BUSINESS_ONLY_NOTICE,
            STATUS_RED if meta.contains_personal_data else TEAL,
        ),
    )
    for text, colour in notices:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row=row, column=1, value=text)
        cell.fill = _fill(colour)
        cell.font = _font(bold=True, colour=_text_colour_on(colour))
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30
        written.append(text)
        row += 2

    contents = ws.cell(
        row=row,
        column=1,
        value="Sheets: " + " · ".join(SHEET_ORDER[1:]),
    )
    contents.font = _font(italic=True)
    written.append(contents.value)
    return written


def _write_funnel(ws: Any, report: SuppressionReport) -> tuple[list[str], list[object]]:
    """Write the suppression funnel, one row per field of `SuppressionReport`.

    Built by iterating the dataclass's fields rather than listing them, so a step added to
    the report cannot quietly go unreported here. A field with no note yet says so in the
    sheet instead of appearing unexplained.

    Returns:
        The headers and every value written.
    """
    rows: list[Sequence[object]] = []
    for field in fields(SuppressionReport):
        value = getattr(report, field.name)
        rendered: object = format_date(value) if isinstance(value, dt.date) else value
        note = _FUNNEL_NOTES.get(
            field.name,
            "Not yet documented. Add a note to _FUNNEL_NOTES in src/grid/export/xlsx.py.",
        )
        rows.append((field.name.replace("_", " ").capitalize(), rendered, note))

    rows.append(
        (
            "Funnel reconciles",
            _yes_no(report.reconciles),
            "Not on panel must equal the call list plus superseded plus both review "
            "routes. No means the numbers on this sheet do not add up and should not be "
            "circulated until they do.",
        )
    )

    written = _write_grid(
        ws,
        headers=list(FUNNEL_HEADERS),
        rows=rows,
        widths=[38, 18, 96],
        notice=FUNNEL_NOTICE,
        notice_colour=DARK_GRAY,
    )
    return list(FUNNEL_HEADERS), written


def _appointment_summary(rows: Sequence[QueueBRow]) -> str:
    """Describe how long these clinics have had a PMCare provider record.

    The median is computed from the population in this very workbook rather than quoted
    from a one-off profile, so the sentence cannot go stale against the data beside it.
    `statistics.median_low` is used so the answer is a real year rather than a half-year.
    """
    years = [row.appointment_date.year for row in rows if row.appointment_date is not None]
    if not years:
        return (
            "No appointment date is present on these rows, which is itself unexpected for "
            "the incumbent master and worth raising."
        )
    return (
        f"{len(years):,} of the {len(rows):,} not-on-panel rows in this workbook carry an "
        f"appointment date, median year {statistics.median_low(years)}. These clinics "
        "already hold a PMCare provider record and are not on the panel today. This is a "
        "re-engagement list, not a list of clinics PMCare has never engaged."
    )


def _write_notes(
    ws: Any, population: QueueBPopulation, meta: WorkbookMeta
) -> tuple[list[str], list[object]]:
    """Write the notes and definitions sheet.

    Returns:
        The headers and every value written.
    """
    not_on_panel = (*population.call_list, *population.review)
    rows: list[Sequence[object]] = [
        (
            "Read this first",
            _appointment_summary(not_on_panel),
            "Known — from the extract",
        ),
        (
            "Chain membership is inferred",
            "Chain grouping comes from a shared base name, not from a company register. It "
            "can mean a real brand, or the same premises keyed twice under two provider "
            "codes. Verify before opening a conversation with 'we already work with your "
            "other branches'.",
            "Inferred",
        ),
        (
            "Panel suppression is a filter",
            "A clinic is suppressed because the incumbent record says it is on the panel, "
            "not because a name looked similar. There is nothing fuzzy to get wrong here.",
            "Known",
        ),
        (
            "Win-back is a different conversation",
            "The Win-back sheet holds clinics that were on the panel and were terminated. "
            "They are not never-engaged prospects; find out why they left first.",
            "Known",
        ),
        (
            "Review rows are not callable",
            "The Review sheet holds rows that are probably already on panel under another "
            "code, or that PMCare's own notes flag. Resolve the reason before contacting.",
            "Known",
        ),
        (
            "Postcode and state are text",
            "Both are written in Excel text format so leading zeros survive. Reformatting "
            "either column as a number turns postcode 05000 into 5000 and state code 00 "
            "into 0, which silently relocates clinics.",
            "Known",
        ),
        ("Status of this workbook", DRAFT_NOTICE, "Known"),
        ("PMCare's role", NEUTRAL_TPA_NOTICE, "Known"),
        (
            "Personal data",
            PERSONAL_DATA_NOTICE if meta.contains_personal_data else BUSINESS_ONLY_NOTICE,
            "Known",
        ),
    ]

    documented: set[str] = set()
    for column in (*QUEUE_B_COLUMNS, *WIN_BACK_COLUMNS, *REVIEW_COLUMNS):
        if column.header in documented:
            continue
        documented.add(column.header)
        meaning, provenance = _COLUMN_NOTES.get(
            column.header,
            (
                "Not yet documented. Add an entry to _COLUMN_NOTES in src/grid/export/xlsx.py.",
                "Unknown",
            ),
        )
        rows.append((column.header, meaning, provenance))

    written = _write_grid(
        ws,
        headers=list(NOTES_HEADERS),
        rows=rows,
        widths=[34, 100, 22],
        notice=NOTES_NOTICE,
        notice_colour=DARK_GRAY,
    )
    return list(NOTES_HEADERS), written


# --------------------------------------------------------------------------------------
# Ranking
# --------------------------------------------------------------------------------------


def _rank(rows: Sequence[QueueBRow], priorities: Mapping[str, Priority]) -> list[_RowView]:
    """Order the call list.

    The ordering itself is `grid.score.priority.priority_key` — band, then the ranking
    signals, then territory, then `provider_code`. It is deliberately not re-implemented
    here: a workbook that ordered rows its own way would be a second, undocumented opinion
    about what "ranked" means, and the two would diverge the first time either changed.

    A row with no priority sorts last rather than being dropped. An unranked clinic is a
    gap in the scoring, and hiding it would hide the gap.

    Args:
        rows: The call list, in any order.
        priorities: Priority by provider code.

    Returns:
        Views in call order, numbered from 1.
    """

    def sort_key(row: QueueBRow) -> tuple[int, tuple[object, ...]]:
        priority = priorities.get(row.provider_code)
        if priority is None:
            return (1, (row.provider_code,))
        return (0, priority_key(row, priority))

    return [
        _RowView(rank=position, row=row, priority=priorities.get(row.provider_code))
        for position, row in enumerate(sorted(rows, key=sort_key), start=1)
    ]


def _plain_views(rows: Sequence[QueueBRow]) -> list[_RowView]:
    """Views for the sheets that are not ranked, preserving the derivation's own order."""
    return [
        _RowView(rank=position, row=row, priority=None)
        for position, row in enumerate(rows, start=1)
    ]


# --------------------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------------------


def write_queue_b_workbook(
    population: QueueBPopulation,
    priorities: Mapping[str, Priority],
    *,
    path: Path,
    meta: WorkbookMeta,
) -> Path:
    """Write the branded Queue B workbook.

    Six sheets: a branded cover, the ranked call list, the win-back list, the review list,
    the suppression funnel so the headline number can be reconciled, and a notes sheet
    stating which values are known and which are inferred.

    The business-only check runs against the assembled workbook *before* the file is saved.
    A workbook that would leak personal data therefore never exists on disk, rather than
    existing briefly and being deleted.

    Args:
        population: The call list, win-back list, review list and funnel counts.
        priorities: Priority by provider code. Rows with no entry are ranked last and
            labelled "Unranked" rather than dropped.
        path: Where to write. Must be inside a ``data/exports`` directory.
        meta: Cover-sheet provenance, including whether personal data is expected.

    Returns:
        The resolved path written.

    Raises:
        ExportPathError: When the path is outside ``data/exports`` or at the repo root.
        PersonalDataInExportError: When a business-only workbook carries personal data.
    """
    destination = _assert_within_exports(path)

    workbook = Workbook()
    cover = workbook.active
    cover.title = SHEET_COVER

    headers: list[str] = []
    values: list[object] = list(_write_cover(cover, meta, population.report))

    queue_views = _rank(population.call_list, priorities)
    queue_sheet = workbook.create_sheet(SHEET_QUEUE_B)
    queue_headers, queue_values = _write_list_sheet(
        queue_sheet,
        QUEUE_B_COLUMNS,
        queue_views,
        notice=QUEUE_B_NOTICE,
        notice_colour=TEAL,
    )
    _colour_band_column(queue_sheet, QUEUE_B_COLUMNS, queue_views)
    headers.extend(queue_headers)
    values.extend(queue_values)

    conflicts = sum(1 for row in population.win_back if row.panel_flag_conflict)
    win_back_notice = WIN_BACK_NOTICE
    if conflicts:
        win_back_notice += (
            f" {conflicts:,} of these are terminated yet still flagged on-panel — a "
            "contradiction in the source that a human must resolve."
        )
    win_back_headers, win_back_values = _write_list_sheet(
        workbook.create_sheet(SHEET_WIN_BACK),
        WIN_BACK_COLUMNS,
        _plain_views(population.win_back),
        notice=win_back_notice,
        notice_colour=STATUS_AMBER,
    )
    headers.extend(win_back_headers)
    values.extend(win_back_values)

    review_headers, review_values = _write_list_sheet(
        workbook.create_sheet(SHEET_REVIEW),
        REVIEW_COLUMNS,
        _plain_views(population.review),
        notice=REVIEW_NOTICE,
        notice_colour=STATUS_RED,
    )
    headers.extend(review_headers)
    values.extend(review_values)

    funnel_headers, funnel_values = _write_funnel(
        workbook.create_sheet(SHEET_FUNNEL), population.report
    )
    headers.extend(funnel_headers)
    values.extend(funnel_values)

    notes_headers, notes_values = _write_notes(workbook.create_sheet(SHEET_NOTES), population, meta)
    headers.extend(notes_headers)
    values.extend(notes_values)

    if not meta.contains_personal_data:
        assert_business_only(headers, values)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)

    log.info(
        "export.queue_b_workbook",
        path=destination.name,
        as_of=meta.as_of.isoformat(),
        call_list=len(population.call_list),
        win_back=len(population.win_back),
        review=len(population.review),
        reconciles=population.report.reconciles,
        contains_personal_data=meta.contains_personal_data,
    )
    return destination
